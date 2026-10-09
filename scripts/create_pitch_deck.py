#!/usr/bin/env python3
"""Create the three-minute Life Off Desk pitch deck as PNG, PDF, and PPTX."""

from pathlib import Path
import json, math, os, statistics, sys

from PIL import Image, ImageDraw, ImageFont, ImageFilter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "pitch"
OUT.mkdir(parents=True, exist_ok=True)
W, H = 1920, 1080

CANVAS = "#F8F6EF"
INK = "#283A31"
GREEN = "#46785B"
MUTED = "#69776D"
SURFACE = "#FFFFFF"
BORDER = "#D9DED4"
DANGER = "#A13D36"
PALE = "#E8EFE6"
AMBER = "#B17620"

FONT_ROOT = Path("/System/Library/Fonts")
FONT_REG = FONT_ROOT / "HelveticaNeue.ttc"
FONT_BOLD = FONT_ROOT / "HelveticaNeue.ttc"


def font(size, bold=False):
    # Faces 0 and 1 in HelveticaNeue.ttc are regular and bold on current macOS.
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REG), size, index=1 if bold else 0)


def canvas():
    im = Image.new("RGB", (W, H), CANVAS)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, W, 14), fill=GREEN)
    return im, d


def text(d, xy, value, size, color=INK, bold=False, anchor=None, spacing=8):
    d.multiline_text(xy, value, font=font(size, bold), fill=color, anchor=anchor, spacing=spacing)


def pill(d, xy, value, fill=PALE, color=GREEN):
    f = font(22, True)
    box = d.textbbox((0, 0), value, font=f)
    tw = box[2] - box[0]
    x, y = xy
    d.rounded_rectangle((x, y, x + tw + 40, y + 48), radius=24, fill=fill)
    d.text((x + 20, y + 11), value, font=f, fill=color)


def footer(d, number, label):
    d.line((90, 1016, 1830, 1016), fill=BORDER, width=2)
    text(d, (90, 1033), "LIFE OFF DESK", 17, MUTED, True)
    text(d, (1830, 1033), f"{number:02d}  ·  {label.upper()}", 17, MUTED, True, "ra")


def rounded_image(path, size, radius=36, crop=True):
    src = Image.open(path).convert("RGB")
    sw, sh = src.size
    tw, th = size
    if crop:
        scale = max(tw / sw, th / sh)
    else:
        scale = min(tw / sw, th / sh)
    nw, nh = round(sw * scale), round(sh * scale)
    src = src.resize((nw, nh), Image.Resampling.LANCZOS)
    left, top = max(0, (nw - tw) // 2), max(0, (nh - th) // 2)
    src = src.crop((left, top, left + tw, top + th))
    mask = Image.new("L", (tw, th), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, tw, th), radius=radius, fill=255)
    out = Image.new("RGB", (tw, th), CANVAS)
    out.paste(src, (0, 0), mask)
    return out


def phone_frame(im, path, xy, size=(330, 710)):
    x, y = xy
    pw, ph = size
    draw = ImageDraw.Draw(im)
    shadow = Image.new("RGBA", im.size, (0, 0, 0, 0))
    sd = ImageDraw.Draw(shadow)
    sd.rounded_rectangle((x + 10, y + 18, x + pw + 10, y + ph + 18), radius=60, fill=(40, 58, 49, 45))
    shadow = shadow.filter(ImageFilter.GaussianBlur(18))
    im.paste(shadow, (0, 0), shadow)
    draw.rounded_rectangle((x, y, x + pw, y + ph), radius=62, fill="#171B19")
    inner = rounded_image(path, (pw - 30, ph - 30), 48)
    im.paste(inner, (x + 15, y + 15))
    draw.rounded_rectangle((x + pw // 2 - 54, y + 17, x + pw // 2 + 54, y + 42), radius=14, fill="#171B19")


def card(d, box, title, body, accent=GREEN):
    x1, y1, x2, y2 = box
    d.rounded_rectangle(box, radius=28, fill=SURFACE, outline=BORDER, width=2)
    d.rectangle((x1, y1, x1 + 10, y2), fill=accent)
    text(d, (x1 + 38, y1 + 28), title, 28, accent, True)
    text(d, (x1 + 38, y1 + 78), body, 23, MUTED, False, spacing=10)


def arrow(d, start, end, color=GREEN, width=5):
    x1, y1 = start
    x2, y2 = end
    d.line((x1, y1, x2 - 18, y2), fill=color, width=width)
    d.polygon(((x2, y2), (x2 - 22, y2 - 13), (x2 - 22, y2 + 13)), fill=color)


def save(im, number, slug):
    path = OUT / f"{number:02d}-{slug}.png"
    im.save(path, quality=95)
    return path


def slide_1():
    im, d = canvas()
    pill(d, (90, 72), "OFFLINE-FIRST · PRIVATE · TAGLISH")
    text(d, (90, 178), "Your world should feel\nbigger than your desk.", 82, INK, True, spacing=4)
    text(d, (95, 395), "Life Off Desk turns a short break into a real,\nlocal adventure — even with no signal.", 34, MUTED, spacing=12)
    text(d, (95, 548), "The problem", 24, GREEN, True)
    text(d, (95, 595), "Desk-bound workers want to step out,\nbut deciding where to go feels like work.", 42, INK, True, spacing=8)
    phone_frame(im, ROOT / "marketing/mockups/map-story.jpg", (1450, 150), (340, 730))
    mascot = rounded_image(ROOT / "LifeOffDesk/Resources/Assets.xcassets/mascot-walking.imageset/mascot-walking.png", (330, 330), 165, False)
    im.paste(mascot, (1125, 650))
    footer(d, 1, "Problem")
    return save(im, 1, "problem")


def slide_2():
    im, d = canvas()
    text(d, (90, 72), "A small nudge. A real walk. A world that remembers.", 55, INK, True)
    text(d, (90, 148), "One offline loop connects intent, movement, discovery, and memory.", 28, MUTED)
    steps = [
        ("01", "ASK", "“Tahimik na park,\n30 mins lang.”"),
        ("02", "WALK", "Track real movement\nand reveal streets."),
        ("03", "REMEMBER", "Save the route,\nphotos, and recap."),
    ]
    xs = [90, 660, 1230]
    for x, (n, h, body) in zip(xs, steps):
        d.rounded_rectangle((x, 260, x + 500, 580), radius=36, fill=SURFACE, outline=BORDER, width=2)
        text(d, (x + 34, 292), n, 24, GREEN, True)
        text(d, (x + 34, 344), h, 39, INK, True)
        text(d, (x + 34, 414), body, 29, MUTED, spacing=10)
        if x < 1230:
            arrow(d, (x + 514, 420), (x + 548, 420))
    d.rounded_rectangle((90, 650, 1175, 914), radius=38, fill=PALE)
    text(d, (135, 696), "What makes it different", 25, GREEN, True)
    text(d, (135, 754), "The AI interprets.\nThe app verifies.", 43, INK, True, spacing=5)
    text(d, (135, 866), "GPS, places, routes, and numbers stay deterministic.", 23, MUTED)
    d.rounded_rectangle((1215, 650, 1830, 914), radius=38, fill="#FFF1EF", outline="#D9A19C", width=2)
    text(d, (1260, 696), "OFFLINE SAFETY HELP", 22, DANGER, True)
    text(d, (1260, 748), "Ask in Taglish.\nGet reviewed steps.", 36, INK, True, spacing=5)
    text(d, (1260, 850), "Urgent language brings Call 911 forward.", 22, MUTED)
    footer(d, 2, "Solution")
    return save(im, 2, "solution")


def slide_3():
    im, d = canvas()
    text(d, (90, 72), "Built first for the person who needs a reason to step out.", 54, INK, True)
    text(d, (90, 148), "Initial customer profile", 25, GREEN, True)
    card(d, (90, 220, 930, 430), "WHO", "Metro Manila desk workers, developers,\nfreelancers, and hybrid teams.")
    card(d, (90, 464, 930, 674), "WHEN", "A low-energy 15–60 minute break\nbetween meetings or after work.")
    card(d, (90, 708, 930, 918), "WHY NOW", "They want novelty without planning overhead,\ndata spend, or sharing location history.")
    phone_frame(im, ROOT / "marketing/mockups/route-story.jpg", (1120, 205), (300, 690))
    phone_frame(im, ROOT / "marketing/mockups/recap-story.jpg", (1465, 205), (300, 690))
    footer(d, 3, "ICP")
    return save(im, 3, "icp")


def slide_4():
    im, d = canvas()
    text(d, (90, 72), "Local AI inside a deterministic walking system.", 56, INK, True)
    text(d, (90, 148), "No backend · no account · no cloud inference", 27, GREEN, True)
    cols = [
        (90, "SENSE", "CoreLocation\nSpeech · Vision", "GPS fixes and optional\nvoice/photo context"),
        (650, "DECIDE", "Qwen3 1.7B\nllama.cpp", "Taglish to constrained JSON\nIDs and intent only"),
        (1210, "VERIFY", "Swift engines\nLocal data", "Filter, rank, route, count,\npersist, and render"),
    ]
    for x, cap, head, body in cols:
        d.rounded_rectangle((x, 255, x + 500, 650), radius=34, fill=SURFACE, outline=BORDER, width=2)
        text(d, (x + 35, 294), cap, 21, GREEN, True)
        text(d, (x + 35, 354), head, 39, INK, True, spacing=8)
        d.line((x + 35, 475, x + 455, 475), fill=BORDER, width=2)
        text(d, (x + 35, 520), body, 27, MUTED, spacing=10)
        if x < 1210:
            arrow(d, (x + 514, 445), (x + 548, 445))
    d.rounded_rectangle((90, 710, 1710, 930), radius=34, fill=INK)
    text(d, (130, 748), "OFFLINE SAFETY ASSISTANT", 20, "#F2B7B1", True)
    text(d, (130, 792), "Taglish question  |  keyword guard  |  reviewed card + local hotlines", 34, SURFACE, True)
    text(d, (130, 850), "AI only routes among 28 bundled topics. It cannot lower a detected emergency or write medical advice.", 23, "#DDE6DF")
    text(d, (130, 887), "The 911 button works without mobile data, but the call still needs cellular signal.", 21, "#F2B7B1")
    footer(d, 4, "Implementation")
    return save(im, 4, "implementation")


def slide_5(metrics):
    im, d = canvas()
    text(d, (90, 72), "The phone proved it offline — and exposed the limit.", 56, INK, True)
    text(d, (90, 148), "iPhone 12 Pro Max · Airplane Mode · 60-case Taglish held-out run · 10 Oct 2026", 25, MUTED)
    stats = [
        ("60 / 60", "schema-valid"),
        ("47 / 60", "intent pass"),
        ("1.59 s", "model load"),
        ("7.38 s", "median generation"),
        ("8.84 s", "p95 generation"),
        ("532 MiB", "peak footprint"),
    ]
    for i, (value, label) in enumerate(stats):
        col, row = i % 3, i // 3
        x, y = 90 + col * 450, 250 + row * 205
        d.rounded_rectangle((x, y, x + 410, y + 165), radius=28, fill=SURFACE, outline=BORDER, width=2)
        text(d, (x + 28, y + 24), value, 45, INK, True)
        text(d, (x + 28, y + 91), label, 23, MUTED)
    d.rounded_rectangle((1460, 250, 1830, 620), radius=34, fill="#FFF1EF", outline="#D9A19C", width=2)
    text(d, (1495, 287), "THERMAL LIMIT", 20, DANGER, True)
    text(d, (1495, 345), "Serious", 35, DANGER, True)
    arrow(d, (1498, 413), (1538, 413), DANGER, 5)
    text(d, (1552, 393), "Critical", 35, DANGER, True)
    text(d, (1495, 476), "during a sustained\n60-request run", 25, MUTED, spacing=8)
    d.rounded_rectangle((90, 700, 1830, 925), radius=34, fill=PALE)
    text(d, (130, 742), "WHAT THIS MEANS", 20, GREEN, True)
    text(d, (130, 791), "Great for short, purposeful requests. Not for continuous generation.", 40, INK, True)
    text(d, (130, 855), "Battery was charging, so energy impact remains unmeasured; outdoor GPS and long-history frame pacing need separate runs.", 24, MUTED)
    footer(d, 5, "Proof + limitation")
    return save(im, 5, "benchmark")


def slide_6():
    im, d = canvas()
    text(d, (90, 72), "From a working local loop to a trusted city companion.", 55, INK, True)
    items = [
        ("NOW", "Prove the walk", "Instrument outdoor GPS, street snapping,\nand large-history map performance."),
        ("NEXT", "Protect the memory", "Add export/restore, compact history indexes,\nand bounded track thinning."),
        ("THEN", "Grow responsibly", "Review accessibility evidence, expand city packs,\nand tune AI against phone heat and quality."),
    ]
    for i, (phase, head, body) in enumerate(items):
        y = 230 + i * 225
        d.ellipse((105, y + 20, 161, y + 76), fill=GREEN)
        text(d, (133, y + 48), str(i + 1), 24, SURFACE, True, "mm")
        if i < 2: d.line((133, y + 77, 133, y + 225), fill="#B8C8BA", width=5)
        text(d, (205, y), phase, 20, GREEN, True)
        text(d, (205, y + 42), head, 37, INK, True)
        text(d, (205, y + 98), body, 26, MUTED, spacing=8)
    mascot = rounded_image(ROOT / "LifeOffDesk/Resources/Assets.xcassets/mascot-discovering.imageset/mascot-discovering.png", (520, 520), 260, False)
    im.paste(mascot, (1260, 360))
    text(d, (1520, 864), "A little walk.\nA bigger world.", 38, INK, True, "mm", spacing=5)
    footer(d, 6, "Future direction")
    return save(im, 6, "future")


def main():
    metrics_path = ROOT / "eval/results/taglish-heldout-iphone12promax-airplane-2026-10-10.json"
    metrics = json.loads(metrics_path.read_text())
    slides = [slide_1(), slide_2(), slide_3(), slide_4(), slide_5(metrics), slide_6()]

    # PDF is directly shareable and preserves the exact rendered layout.
    ims = [Image.open(p).convert("RGB") for p in slides]
    ims[0].save(OUT / "life-off-desk-3-minute-pitch.pdf", save_all=True, append_images=ims[1:], resolution=150)

    # PPTX uses full-slide images so Keynote, PowerPoint, and web viewers render identically.
    sys.path.insert(0, "/tmp/lod-slide-deps")
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(OUT / "life-off-desk-3-minute-pitch.pptx")
    print("Created", len(slides), "slides in", OUT)


if __name__ == "__main__":
    main()
